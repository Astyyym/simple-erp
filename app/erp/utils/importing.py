from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from copy import copy

from flask import send_file
from openpyxl import Workbook, load_workbook

MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_DISPLAYED_ERRORS = 100


class ImportFileError(ValueError):
    pass


@dataclass
class ImportResult:
    added: int = 0
    skipped: int = 0
    errors: list[str] = field(default_factory=list)
    error_count: int = 0

    def add_error(self, message: str) -> None:
        self.error_count += 1
        if len(self.errors) < MAX_DISPLAYED_ERRORS:
            self.errors.append(message)


def excel_template(headers: list[str], filename: str):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "导入模板"
    sheet.append(headers)
    for cell in sheet[1]:
        font = copy(cell.font)
        font.bold = True
        cell.font = font
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        max_age=0,
    )


def read_upload(upload, expected_headers: list[str] | tuple[list[str], ...]) -> list[tuple[int, list[object]]]:
    _headers, rows = read_upload_detailed(upload, expected_headers)
    return rows


def read_upload_detailed(
    upload, expected_headers: list[str] | tuple[list[str], ...]
) -> tuple[list[str], list[tuple[int, list[object]]]]:
    """与 read_upload 相同的校验，但额外返回表头，供按列名映射使用。"""
    filename = (upload.filename or "").strip()
    suffix = Path(filename).suffix.lower()
    if suffix not in {".xlsx", ".csv"}:
        raise ImportFileError("仅支持 .xlsx 或 .csv 文件")
    data = upload.stream.read(MAX_IMPORT_BYTES + 1)
    if len(data) > MAX_IMPORT_BYTES:
        raise ImportFileError("文件不能超过5MB")
    if not data:
        raise ImportFileError("文件为空")
    try:
        rows = _xlsx_rows(data) if suffix == ".xlsx" else _csv_rows(data)
    except ImportFileError:
        raise
    except Exception as exc:
        raise ImportFileError("文件无法解析，请确认文件未损坏") from exc
    if not rows:
        raise ImportFileError("文件为空")
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    accepted_headers = (expected_headers,) if expected_headers and isinstance(expected_headers[0], str) else expected_headers
    if headers not in accepted_headers:
        expected = " 或 ".join("、".join(candidate) for candidate in accepted_headers)
        raise ImportFileError(f"表头错误，应为：{expected}")
    if len(rows) == 1:
        raise ImportFileError("没有可导入的数据行")
    return headers, [(number, list(row)) for number, row in enumerate(rows[1:], start=2)]


def _xlsx_rows(data: bytes) -> list[tuple[object, ...]]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.active
    if sheet is None:
        return []
    return list(sheet.iter_rows(values_only=True))


def _csv_rows(data: bytes) -> list[tuple[str, ...]]:
    text = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ImportFileError("CSV编码无法识别，请使用UTF-8或GB18030")
    return [tuple(row) for row in csv.reader(io.StringIO(text))]


def normalized_name(value: object) -> str:
    return "" if value is None else str(value).strip()


def price_to_cents(value: object) -> int:
    if value is None or str(value).strip() == "":
        raise ValueError("价格不能为空")
    try:
        amount = Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        raise ValueError("价格格式无效") from None
    if not amount.is_finite():
        raise ValueError("价格格式无效")
    if amount < 0:
        raise ValueError("价格不能为负数")
    return int((amount * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
