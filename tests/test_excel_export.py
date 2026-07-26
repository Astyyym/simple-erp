from __future__ import annotations

import io
from uuid import uuid4

from openpyxl import load_workbook

from erp import create_app
from erp.db import get_db, init_db
from erp.services.accounting import create_customer, create_order_from_typed_rows, create_product
from erp.utils.exporting import (
    customer_export_headers,
    order_line_export_headers,
    product_export_headers,
)
from erp.utils.money import yuan_to_cents


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _sheet_rows(response):
    assert response.status_code == 200
    assert XLSX in response.content_type
    assert "attachment" in response.headers.get("Content-Disposition", "")
    workbook = load_workbook(io.BytesIO(response.data), read_only=True, data_only=True)
    sheet = workbook.active
    return list(sheet.iter_rows(values_only=True))


def test_excel_export_endpoints_and_headers():
    init_db()
    suffix = uuid4().hex[:8]
    customer_name = f"导出客户{suffix}"
    product_name = f"导出商品{suffix}"
    create_customer(customer_name, phone="13800000000", address="杭州", opening_balance_cents=yuan_to_cents("12.50"))
    create_product(product_name, "DN50", "个", yuan_to_cents("66.00"))
    customer_id = create_customer(f"单据客户{suffix}")
    create_order_from_typed_rows(
        customer_id,
        f"EXP-SALE-{suffix}",
        [
            {"product_name": "阀门A", "spec": "DN65", "unit": "个", "unit_price_yuan": "100.00", "quantity": "2"},
            {"product_name": "阀门B", "spec": "", "unit": "套", "unit_price_yuan": "50.50", "quantity": "1"},
        ],
        status="saved",
        order_date="2026-07-10",
        notes="销售备注",
        order_type="sale",
    )
    create_order_from_typed_rows(
        customer_id,
        f"EXP-RET-{suffix}",
        [{"product_name": "退货品", "unit": "个", "unit_price_yuan": "20.00", "quantity": "1"}],
        status="saved",
        order_date="2026-07-11",
        notes="退货备注",
        order_type="return",
    )
    # Soft-deleted rows must not appear in master-data exports.
    with get_db() as conn:
        conn.execute(
            "UPDATE customers SET deleted_at=CURRENT_TIMESTAMP WHERE name=?",
            (customer_name,),
        )
        conn.execute(
            "UPDATE products SET deleted_at=CURRENT_TIMESTAMP WHERE name=?",
            (product_name,),
        )
        alive_customer = f"在用客户{suffix}"
        conn.execute(
            "INSERT INTO customers(name, phone, address, opening_balance_cents, notes) VALUES (?, '139', '宁波', 250, '备注X')",
            (alive_customer,),
        )
        conn.execute(
            "INSERT INTO products(name, spec, unit, default_price_cents, notes, usage_count) VALUES (?, 'S1', '箱', 1234, '商品备注', 3)",
            (f"在用商品{suffix}",),
        )

    client = create_app().test_client()

    customers_page = client.get("/customers/").get_data(as_text=True)
    products_page = client.get("/products/").get_data(as_text=True)
    orders_page = client.get("/orders/").get_data(as_text=True)
    assert "导出Excel" in customers_page and "/customers/export.xlsx" in customers_page
    assert "导出Excel" in products_page and "/products/export.xlsx" in products_page
    assert "导出销售单Excel" in orders_page and "/orders/export/sales.xlsx" in orders_page
    assert "导出退货单Excel" in orders_page and "/orders/export/returns.xlsx" in orders_page
    assert "导出汇总表" in orders_page and "/orders/summary_pdf" in orders_page

    customer_rows = _sheet_rows(client.get("/customers/export.xlsx"))
    assert list(customer_rows[0]) == customer_export_headers()
    customer_names = {row[0] for row in customer_rows[1:]}
    assert alive_customer in customer_names
    assert customer_name not in customer_names
    alive_row = next(row for row in customer_rows[1:] if row[0] == alive_customer)
    assert alive_row[1] == "139"
    assert alive_row[2] == "宁波"
    assert str(alive_row[3]) == "2.50"
    assert alive_row[4] == "备注X"

    product_rows = _sheet_rows(client.get("/products/export.xlsx"))
    assert list(product_rows[0]) == product_export_headers()
    product_names = {row[0] for row in product_rows[1:]}
    assert f"在用商品{suffix}" in product_names
    assert product_name not in product_names
    product_row = next(row for row in product_rows[1:] if row[0] == f"在用商品{suffix}")
    assert product_row[1] == "S1"
    assert product_row[2] == "箱"
    assert str(product_row[3]) == "12.34"
    assert product_row[4] == "商品备注"
    assert int(product_row[5]) == 3

    sale_rows = _sheet_rows(
        client.get(
            "/orders/export/sales.xlsx",
            query_string={
                "customer": f"单据客户{suffix}",
                "date_mode": "range",
                "start_date": "2026-07-01",
                "end_date": "2026-07-31",
            },
        )
    )
    assert list(sale_rows[0]) == order_line_export_headers()
    assert len(sale_rows) == 3  # header + 2 detail lines
    assert {row[0] for row in sale_rows[1:]} == {f"EXP-SALE-{suffix}"}
    assert all(row[3] == "销售单" for row in sale_rows[1:])
    unit_prices = sorted(str(row[9]) for row in sale_rows[1:])
    amounts = sorted(str(row[10]) for row in sale_rows[1:])
    assert unit_prices == ["100.00", "50.50"]
    assert amounts == ["200.00", "50.50"]  # 2*100 and 1*50.50
    totals = {str(row[12]) for row in sale_rows[1:]}
    assert totals == {"250.50"}
    assert sale_rows[1][11] == "销售备注" or sale_rows[2][11] == "销售备注"

    return_rows = _sheet_rows(
        client.get(
            "/orders/export/returns.xlsx",
            query_string={
                "customer": f"单据客户{suffix}",
                "date_mode": "range",
                "start_date": "2026-07-01",
                "end_date": "2026-07-31",
            },
        )
    )
    assert list(return_rows[0]) == order_line_export_headers()
    assert len(return_rows) == 2
    assert return_rows[1][0] == f"EXP-RET-{suffix}"
    assert return_rows[1][3] == "退货单"
    assert str(return_rows[1][9]) == "20.00"
    assert str(return_rows[1][10]) == "20.00"
    # Return order total is stored negative; export uses cents_to_yuan as-is.
    assert str(return_rows[1][12]) == "-20.00"
    assert return_rows[1][11] == "退货备注"


def test_excel_export_open_without_login(monkeypatch):
    monkeypatch.delenv("ERP_DISABLE_AUTH", raising=False)
    init_db()
    client = create_app().test_client()
    for url in (
        "/customers/export.xlsx",
        "/products/export.xlsx",
        "/orders/export/sales.xlsx",
        "/orders/export/returns.xlsx",
    ):
        resp = client.get(url, follow_redirects=False)
        assert resp.status_code == 200, url
        assert resp.mimetype in (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/octet-stream",
        ) or (resp.data[:2] == b"PK")
