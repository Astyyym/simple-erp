from __future__ import annotations

import io
import re
import uuid

from openpyxl import load_workbook

from erp import create_app
from erp.db import get_db, init_db


def _name(prefix: str) -> str:
    return f"{prefix}-{uuid.uuid4().hex[:10]}"


def _upload(client, url: str, filename: str, content: bytes, *, same_name: str = "skip"):
    return client.post(
        url,
        data={"file": (io.BytesIO(content), filename), "same_name": same_name},
        content_type="multipart/form-data",
        follow_redirects=True,
    )


def _preview_json(html: str) -> str:
    m = re.search(r'name="preview_json" value="([^"]*)"', html)
    assert m, "预检页必须携带 preview_json 隐藏域"
    import html as html_module

    return html_module.unescape(m.group(1))


def _confirm(client, confirm_url: str, preview_json: str, *, same_name: str = "skip"):
    return client.post(
        confirm_url,
        data={"preview_json": preview_json, "same_name": same_name},
        follow_redirects=True,
    )


def test_product_and_customer_excel_templates_have_extended_headers():
    client = create_app().test_client()

    product_response = client.get("/products/import/template")
    customer_response = client.get("/customers/import/template")

    assert product_response.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in product_response.content_type
    assert "attachment" in product_response.headers["Content-Disposition"]
    product_sheet = load_workbook(io.BytesIO(product_response.data), read_only=True).active
    # 决策变更（v14）：商品身份 = 单一名称，模板不再有「型号」列（型号直接写进商品名称）。
    # 决策变更（2026-10-15）：默认价下线，模板不再有「默认价」列（售价唯一来源 = 开单学习）。
    # 决策变更（2026-10-16 ⑥）：库存系统下线，商品模板去掉「期初数量 / 期初成本」两列。
    assert list(next(product_sheet.iter_rows(values_only=True))) == [
        "商品名称", "品牌", "单位", "备注"
    ]

    assert customer_response.status_code == 200
    customer_sheet = load_workbook(io.BytesIO(customer_response.data), read_only=True).active
    assert list(next(customer_sheet.iter_rows(values_only=True))) == [
        "客户名称", "电话", "地址", "期初余额（元）", "备注"
    ]


def test_management_pages_show_chinese_import_controls():
    client = create_app().test_client()
    products = client.get("/products/").get_data(as_text=True)
    customers = client.get("/customers/").get_data(as_text=True)

    assert "/products/import" in products and "下载Excel模板" in products and 'accept=".xlsx,.csv"' in products
    assert "/customers/import" in customers and "下载Excel模板" in customers and 'accept=".xlsx,.csv"' in customers
    assert "预检并导入" in products and "预检并导入" in customers


def test_product_import_preview_then_confirm_writes_rows():
    """导入不再读写售价：文件里的价格列被忽略，其余列照常落库。

    旧断言「默认价留空按 0 落库 / 负价报错」随字段下线失效；2026-10-16（决策 ⑥ + schema v15）
    后商品档案的「默认价」列本身也已随 schema v15 删除，改为钉住新契约：
    价格列不解析也不校验（留空、负数都照样导入），文件里的其它列照常落库；
    只有「商品名称」仍是必填校验项。
    """
    init_db()
    existing = _name("已有商品")
    fresh = _name("新增商品")
    with get_db() as conn:
        conn.execute("INSERT INTO products(name, unit) VALUES (?, '箱')", (existing,))
    csv_data = f"商品名称,价格\n{fresh},12.34\n{existing},88\n{fresh},20\n空价商品,\n负价商品,-1\n,10\n".encode("utf-8-sig")

    client = create_app().test_client()
    preview = _upload(client, "/products/import", "products.csv", csv_data)
    html = preview.get_data(as_text=True)

    # 预检：未确认不得落库。
    assert preview.status_code == 200
    assert "将新增：3" in html and "跳过：2" in html and "错误：1" in html
    # 价格列已下线：不再有价格相关的校验错误。
    assert "价格不能为空" not in html and "价格不能为负数" not in html
    assert "第7行" in html and "商品名称不能为空" in html
    with get_db() as conn:
        assert conn.execute("SELECT 1 FROM products WHERE name=?", (fresh,)).fetchone() is None

    # 确认后才落库。
    confirmed = _confirm(client, "/products/import/confirm", _preview_json(html))
    assert confirmed.status_code == 200
    assert "新增：3" in confirmed.get_data(as_text=True)
    with get_db() as conn:
        row = conn.execute("SELECT * FROM products WHERE name=?", (fresh,)).fetchone()
        old = conn.execute("SELECT * FROM products WHERE name=?", (existing,)).fetchone()
        blank_price = conn.execute("SELECT * FROM products WHERE name='空价商品'").fetchone()
    assert row["unit"] == "个"
    # 价格列已下线：文件里的价格既不入库、也不影响其它列（单位不写则回落「个」）。
    assert old["unit"] == "箱"
    assert blank_price["origin"] == "import"


def test_customer_import_preview_counts_soft_deleted_and_file_duplicates():
    init_db()
    deleted = _name("软删客户")
    fresh = _name("新增客户")
    with get_db() as conn:
        conn.execute("INSERT INTO customers(name, deleted_at) VALUES (?, CURRENT_TIMESTAMP)", (deleted,))

    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["客户名称"])
    sheet.append([fresh])
    sheet.append([deleted])
    sheet.append([fresh])
    output = io.BytesIO()
    workbook.save(output)

    client = create_app().test_client()
    preview = _upload(client, "/customers/import", "客户.xlsx", output.getvalue())
    html = preview.get_data(as_text=True)
    assert preview.status_code == 200
    assert "将新增：1" in html and "跳过：1" in html
    # B-8：软删占名明确提示，不静默跳过。
    assert "回收站" in html

    confirmed = _confirm(client, "/customers/import/confirm", _preview_json(html))
    assert confirmed.status_code == 200
    with get_db() as conn:
        created = conn.execute("SELECT * FROM customers WHERE name=?", (fresh,)).fetchone()
        still_deleted = conn.execute("SELECT * FROM customers WHERE name=?", (deleted,)).fetchone()
    assert created["origin"] == "import"
    assert still_deleted["deleted_at"] is not None


def test_import_rejects_bad_suffix_header_empty_file_and_oversize_without_500():
    init_db()
    keep_name = _name("保留商品")
    with get_db() as conn:
        conn.execute("INSERT INTO products(name, unit) VALUES (?, '个')", (keep_name,))
    client = create_app().test_client()
    cases = [
        ("/products/import", "bad.txt", b"x", "仅支持 .xlsx 或 .csv 文件"),
        ("/products/import", "empty.csv", b"", "文件为空"),
        ("/products/import", "header-only.csv", "商品名称,价格\n".encode(), "没有可导入的数据行"),
        ("/products/import", "bad.csv", "名称,金额\n商品,1\n".encode(), "表头错误"),
        ("/customers/import", "large.csv", b"a" * (5 * 1024 * 1024 + 1), "文件不能超过5MB"),
    ]
    for url, filename, content, expected in cases:
        response = _upload(client, url, filename, content)
        html = response.get_data(as_text=True)
        assert response.status_code == 400
        assert expected in html
        if url.startswith("/products/"):
            assert keep_name in html
            assert "下载Excel模板" in html


def test_customer_csv_import_accepts_gb18030():
    init_db()
    fresh = _name("国标客户")
    client = create_app().test_client()
    preview = _upload(
        client,
        "/customers/import",
        "客户.csv",
        f"客户名称\n{fresh}\n".encode("gb18030"),
    )
    html = preview.get_data(as_text=True)
    assert preview.status_code == 200
    confirmed = _confirm(client, "/customers/import/confirm", _preview_json(html))
    assert "新增：1" in confirmed.get_data(as_text=True)
    with get_db() as conn:
        assert conn.execute("SELECT id FROM customers WHERE name=?", (fresh,)).fetchone() is not None


def test_product_import_rejects_spec_header_and_still_supports_two_column_header():
    """决策变更（v14）：商品身份 = 单一名称，`spec`（型号）已并入名称。

    原「同名不同型号」用例依赖已删除的型号列，改为钉住新契约：
    - 含「型号」列的旧文件必须**明确报错（表头错误）**，不静默丢型号、也不落库；
    - 旧的两列模板「商品名称,价格」仍可导入（单位回落「个」），旧库不被阻塞。
    """
    init_db()
    name = _name("同名型号商品")
    deleted = _name("回收站商品")
    with get_db() as conn:
        conn.execute("INSERT INTO products(name, unit) VALUES (?, '个')", (name,))
        conn.execute("INSERT INTO products(name, unit, deleted_at) VALUES (?, '个', CURRENT_TIMESTAMP)", (deleted,))

    # 含「型号」列的旧文件：整表头精确匹配挡下，明确报错。
    legacy_csv = (
        f"商品名称,型号,价格\n{name}, DN20 ,99\n{name},DN25,20\n{deleted},DN25,50\n"
    ).encode("utf-8-sig")
    client = create_app().test_client()
    rejected = _upload(client, "/products/import", "products.csv", legacy_csv)
    rejected_html = rejected.get_data(as_text=True)
    assert rejected.status_code == 400
    assert "表头错误" in rejected_html
    # 报错即零写入：既有商品价不动，回收站记录不动，也不新增。
    with get_db() as conn:
        assert conn.execute("SELECT COUNT(*) FROM products").fetchone()[0] == 2
        assert conn.execute("SELECT unit FROM products WHERE name=?", (name,)).fetchone()[0] == "个"
        assert conn.execute("SELECT deleted_at FROM products WHERE name=?", (deleted,)).fetchone()[0] is not None

    # 旧两列模板仍可导入（单位回落「个」）。
    old_name = _name("旧格式商品")
    old_csv = f"商品名称,价格\n{old_name},8.88\n".encode("utf-8-sig")
    old_preview = _upload(client, "/products/import", "old.csv", old_csv)
    old_html = old_preview.get_data(as_text=True)
    assert old_preview.status_code == 200
    _confirm(client, "/products/import/confirm", _preview_json(old_html))
    with get_db() as conn:
        old_row = conn.execute("SELECT unit FROM products WHERE name=?", (old_name,)).fetchone()
    # 旧两列模板「商品名称,价格」仍可导入：价格列已下线（不再解析），单位回落「个」。
    assert old_row["unit"] == "个"


def test_incomplete_import_rows_import_anyway_and_get_flagged_todo():
    """字段不完整的导入行：照样入库，但三处都提醒「待补全」（用户 2026-10-14 第二轮要求）。

    三处 = 预检页提前提醒 + 导入结果摘要 + 列表徽标与「待补全」筛选。
    口径见 `services/master_data_quality.py`：来源 order/import 且**缺品牌**
    （2026-10-15 默认价下线后，商品侧只剩品牌这一项）。
    """
    init_db()
    client = create_app().test_client()
    no_brand = _name("导入缺品牌商品")
    with_brand = _name("导入有品牌商品")
    product_csv = (
        "商品名称,品牌,单位,备注\n"
        f"{no_brand},,件,验证\n"
        f"{with_brand},金枪,件,验证\n"
    ).encode("utf-8-sig")
    preview = _upload(client, "/products/import", "products.csv", product_csv)
    preview_html = preview.get_data(as_text=True)
    assert preview.status_code == 200
    assert "将新增：2" in preview_html and "错误：0" in preview_html
    # ① 预检页：条数 + 该行缺什么（只有缺品牌那一行算不完整）。
    assert "有 1 条新增行字段不完整" in preview_html
    assert "缺品牌" in preview_html
    assert "导入后会标记为「待补全」" in preview_html

    confirmed = _confirm(client, "/products/import/confirm", _preview_json(preview_html))
    done_html = confirmed.get_data(as_text=True)
    # ② 导入结果摘要。
    assert "新增：2" in done_html
    assert "其中 <strong>1</strong> 条字段不完整" in done_html

    # ③ 列表徽标 + 筛选（口径与徽标同源）。
    with get_db() as conn:
        brand_row = conn.execute("SELECT * FROM products WHERE name=?", (with_brand,)).fetchone()
        assert brand_row["origin"] == "import"
        assert brand_row["unit"] == "件"
    todo_html = client.get("/products/?quality=todo").get_data(as_text=True)
    assert no_brand in todo_html and with_brand not in todo_html
    body = client.get("/products/").get_data(as_text=True).split('id="productsResults"', 1)[1]
    for name, flagged in ((no_brand, True), (with_brand, False)):
        row = [line for line in body.split("<tr ") if name in line]
        assert row, f"{name} 应在列表里"
        assert ("待补全</span>" in row[0]) is flagged, f"{name} 待补全徽标应为 {flagged}"

    # 补齐品牌后自动移出（徽标、筛选、chip 计数三处同时生效）。
    from erp.services.accounting import update_product_record

    with get_db() as conn:
        product_id = conn.execute("SELECT id FROM products WHERE name=?", (no_brand,)).fetchone()["id"]
    update_product_record(product_id, no_brand, "件", brand="金枪")
    assert no_brand not in client.get("/products/?quality=todo").get_data(as_text=True)


def test_incomplete_customer_import_lands_in_todo():
    init_db()
    client = create_app().test_client()
    name = _name("导入缺字段客户")
    csv_data = f"客户名称,电话,地址,期初余额（元）,备注\n{name},,,0,验证\n".encode("utf-8-sig")
    preview = _upload(client, "/customers/import", "customers.csv", csv_data)
    preview_html = preview.get_data(as_text=True)
    assert "有 1 条新增行字段不完整" in preview_html
    assert "缺联系方式" in preview_html

    confirmed = _confirm(client, "/customers/import/confirm", _preview_json(preview_html))
    done_html = confirmed.get_data(as_text=True)
    assert "新增：1" in done_html
    assert "其中 <strong>1</strong> 条字段不完整" in done_html
    assert name in client.get("/customers/?quality=todo").get_data(as_text=True)
    body = client.get("/customers/").get_data(as_text=True).split('id="customersResults"', 1)[1]
    row = [line for line in body.split("<tr>") if name in line]
    assert row and "待补全</span>" in row[0]


def test_templates_ship_a_fill_in_guide_sheet_and_keep_the_header_row_intact():
    """模板第二个 sheet「填写说明」写明「留空 → 待补全」；第一行仍是纯表头。

    第一行不能塞说明：`read_upload_detailed` 按整表头精确匹配，多一行会被当数据行或报「表头错误」。
    """
    init_db()
    client = create_app().test_client()
    cases = (
        ("/products/import/template", ["商品名称", "品牌", "单位", "备注"]),
        ("/customers/import/template", ["客户名称", "电话", "地址", "期初余额（元）", "备注"]),
    )
    for url, headers in cases:
        response = client.get(url)
        assert response.status_code == 200
        workbook = load_workbook(io.BytesIO(response.data))
        assert list(next(workbook.active.iter_rows(values_only=True))) == headers, url
        assert "填写说明" in workbook.sheetnames, url
        notes = "\n".join(str(cell[0].value or "") for cell in workbook["填写说明"].iter_rows(min_col=1, max_col=1))
        assert "待补全" in notes, url
        assert "留空" in notes, url
