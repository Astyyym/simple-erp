"""E-2：商品品牌（可选）字段。

计划条目（`开发短计划/2026-10-08_旧数据迁移承接与备份恢复_执行计划.md` §五 Batch E）：
products 加**可空** `brand` 列；导入可空；列表/编辑页可填可显示；**不做唯一键、不做必填、不阻塞迁移**。

本文件锁四件事：
1. schema 迁移幂等，旧库（无 brand 列）加列后旧数据留空、**不猜**；
2. 品牌可空、写入归一成 NULL，且**不参与业务唯一键**（同名同型号不同品牌仍是同一件商品）；
3. 导入新旧表头都能走通（无品牌列 → 品牌留空）；
4. 页面（新建/编辑/列表）与导出的品牌列真的落地。
"""

from __future__ import annotations

import html as html_module
import io
import re

from openpyxl import load_workbook

from erp import create_app
from erp.db import get_db, init_db
from erp.services.accounting import create_product, update_product_record

# v3.2.0（E-2 之前）的 products 表：没有 brand 列，用来模拟旧库。
LEGACY_PRODUCTS_DDL = """
CREATE TABLE products (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT NOT NULL,
    spec TEXT DEFAULT '',
    unit TEXT NOT NULL,
    default_price_cents INTEGER NOT NULL CHECK(default_price_cents >= 0),
    notes TEXT DEFAULT '',
    is_active INTEGER NOT NULL DEFAULT 1,
    usage_count INTEGER NOT NULL DEFAULT 0,
    pinyin_initials TEXT DEFAULT '',
    origin TEXT,
    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
    updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
    deleted_at TEXT,
    delete_reason TEXT DEFAULT '',
    image_path TEXT NOT NULL DEFAULT '',
    safety_stock_3dp INTEGER NOT NULL DEFAULT 0 CHECK(safety_stock_3dp >= 0)
)
"""


def _columns(table: str) -> list[str]:
    with get_db() as conn:
        return [row["name"] for row in conn.execute(f"PRAGMA table_info({table})")]


def _brand_of(product_id: int):
    with get_db() as conn:
        return conn.execute("SELECT brand FROM products WHERE id=?", (product_id,)).fetchone()["brand"]


def _preview_json(resp) -> str:
    m = re.search(r'name="preview_json" value="([^"]*)"', resp.get_data(as_text=True))
    assert m, "预检页必须携带 preview_json 隐藏域"
    return html_module.unescape(m.group(1))


def _upload(client, url, filename, content, *, same_name="skip"):
    return client.post(
        url,
        data={"file": (io.BytesIO(content), filename), "same_name": same_name},
        content_type="multipart/form-data",
    )


# --------------------------- schema 迁移 ---------------------------

def test_brand_column_migration_is_idempotent_and_keeps_legacy_rows_blank():
    init_db()
    with get_db() as conn:
        # 用旧 schema 重建 products，塞一条「旧软件导出的商品」。
        conn.execute("DROP TABLE products")
        conn.execute(LEGACY_PRODUCTS_DDL)
        conn.execute(
            "INSERT INTO products(name, spec, unit, default_price_cents, origin) VALUES ('旧库商品', 'DN20', '个', 1234, NULL)"
        )
    assert "brand" not in _columns("products")

    init_db()  # 第一次迁移：加列
    assert "brand" in _columns("products")
    with get_db() as conn:
        legacy = conn.execute("SELECT brand FROM products WHERE name='旧库商品'").fetchone()
    assert legacy["brand"] is None, "旧数据一律留空，不根据名称等间接信息猜品牌"

    init_db()  # 第二次迁移：必须幂等，不重复加列也不报错
    assert _columns("products").count("brand") == 1
    with get_db() as conn:
        assert conn.execute("PRAGMA integrity_check").fetchone()[0] == "ok"


def test_brand_is_not_part_of_the_business_identity_key():
    init_db()
    first = create_product("阀门", "DN20", "个", 1000, brand="沪工")
    assert _brand_of(first) == "沪工"
    # 品牌不进唯一键：同名同型号换品牌仍是同一件商品，必须被唯一约束挡住。
    try:
        create_product("阀门", "DN20", "个", 1000, brand="远大")
    except ValueError as exc:
        assert "已存在" in str(exc)
    else:
        raise AssertionError("品牌不得参与业务唯一键，同名同型号应被拒绝")
    # 同名不同型号可以各带品牌。
    second = create_product("阀门", "DN25", "个", 1000, brand="远大")
    assert _brand_of(second) == "远大"


# --------------------------- 服务层写入 ---------------------------

def test_blank_brand_is_stored_as_null_not_empty_string():
    init_db()
    for value in ("", "   ", None):
        pid = create_product(f"空品牌{value!r}", "", "个", 100, brand=value)
        assert _brand_of(pid) is None


def test_update_product_sets_and_clears_brand_but_omitted_brand_is_preserved():
    init_db()
    pid = create_product("编辑品牌商品", "S1", "个", 2000, brand="沪工")

    # 显式传品牌 → 覆盖
    update_product_record(pid, "编辑品牌商品", "S1", "个", 2000, brand="远大")
    assert _brand_of(pid) == "远大"

    # 显式传空 → 清空（用户主动删掉品牌）
    update_product_record(pid, "编辑品牌商品", "S1", "个", 2000, brand="")
    assert _brand_of(pid) is None

    # 不传 brand（旧调用点/局部更新）→ 保持原值，不误清
    update_product_record(pid, "编辑品牌商品", "S1", "个", 2000, brand="沪工")
    update_product_record(pid, "编辑品牌商品", "S1", "个", 2100)
    assert _brand_of(pid) == "沪工"


# --------------------------- 导入 ---------------------------

def test_product_import_template_has_optional_brand_column():
    client = create_app().test_client()
    response = client.get("/products/import/template")
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.data), read_only=True).active
    assert list(next(sheet.iter_rows(values_only=True))) == [
        "商品名称", "型号", "品牌", "单位", "默认价", "期初数量", "期初成本（元）", "备注"
    ]


def test_product_import_with_brand_writes_it_and_blank_stays_null():
    init_db()
    csv = (
        "商品名称,型号,品牌,单位,默认价,期初数量,期初成本（元）,备注\n"
        "有品牌商品,4kg,沪工,个,120,,\n"
        "空品牌商品,8kg,,个,220,,\n"
    ).encode("utf-8-sig")
    client = create_app().test_client()
    preview = _upload(client, "/products/import", "p.csv", csv)
    assert preview.status_code == 200
    client.post("/products/import/confirm", data={"preview_json": _preview_json(preview), "same_name": "skip"}, follow_redirects=True)
    with get_db() as conn:
        rows = {row["name"]: row["brand"] for row in conn.execute("SELECT name, brand FROM products")}
    assert rows["有品牌商品"] == "沪工"
    assert rows["空品牌商品"] is None


def test_product_import_still_accepts_pre_brand_header():
    """E-2 之前的表头（无品牌列）必须继续可导入，品牌留空——不阻塞迁移。"""
    init_db()
    csv = "商品名称,型号,单位,默认价,期初数量,期初成本（元）,备注\n旧表头商品,4kg,个,120,,\n".encode("utf-8-sig")
    client = create_app().test_client()
    preview = _upload(client, "/products/import", "old.csv", csv)
    assert preview.status_code == 200
    client.post("/products/import/confirm", data={"preview_json": _preview_json(preview), "same_name": "skip"}, follow_redirects=True)
    with get_db() as conn:
        row = conn.execute("SELECT brand, unit FROM products WHERE name='旧表头商品'").fetchone()
    assert row["brand"] is None
    assert row["unit"] == "个"


def test_product_import_update_branch_can_fill_brand():
    init_db()
    create_product("补品牌商品", "S1", "个", 500)
    csv = "商品名称,型号,品牌,单位,默认价,期初数量,期初成本（元）,备注\n补品牌商品,S1,沪工,个,600,,\n".encode("utf-8-sig")
    client = create_app().test_client()
    preview = _upload(client, "/products/import", "p.csv", csv, same_name="update")
    assert "将更新：1" in preview.get_data(as_text=True)
    client.post("/products/import/confirm", data={"preview_json": _preview_json(preview), "same_name": "update"}, follow_redirects=True)
    with get_db() as conn:
        row = conn.execute("SELECT brand, default_price_cents FROM products WHERE name='补品牌商品'").fetchone()
    assert row["brand"] == "沪工"
    assert row["default_price_cents"] == 60000


# --------------------------- 页面与导出 ---------------------------

def test_product_pages_show_optional_brand_controls_and_value():
    init_db()
    pid = create_product("页面品牌商品", "S1", "个", 2000, brand="沪工")
    client = create_app().test_client()

    list_html = client.get("/products/").get_data(as_text=True)
    assert "<th>品牌</th>" in list_html
    assert 'name="brand" placeholder="品牌(可选)"' in list_html
    assert "沪工" in list_html

    edit_html = client.get(f"/products/{pid}/edit").get_data(as_text=True)
    assert 'name="brand" value="沪工"' in edit_html

    # 编辑页提交空品牌 → 真的清空（不是「不填就保留」）。
    client.post(f"/products/{pid}/edit", data={
        "name": "页面品牌商品", "spec": "S1", "brand": "", "unit": "个",
        "default_price": "20", "safety_stock": "0",
    })
    assert _brand_of(pid) is None


def test_product_export_includes_brand_column():
    init_db()
    create_product("导出品牌商品", "S1", "个", 1234, brand="沪工")
    client = create_app().test_client()
    response = client.get("/products/export.xlsx")
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.data), read_only=True, data_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    assert list(rows[0])[:4] == ["商品名称", "型号", "品牌", "单位"]
    row = next(r for r in rows[1:] if r[0] == "导出品牌商品")
    assert row[2] == "沪工"
