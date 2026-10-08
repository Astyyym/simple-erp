"""Real isolated product routes: independent enable and current stock status."""
from datetime import date
from html.parser import HTMLParser

import pytest

from erp import create_app
from erp.db import get_db
from erp.services.accounting import create_product
from erp.services.inventory import initialize_product


class ProductHTML(HTMLParser):
    def __init__(self, html):
        super().__init__()
        self.rows, self.definitions = [], {}
        self.row = self.cell = self.definition = None
        self.term = None
        self.feed(html)

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self.row = []
        elif tag in {"th", "td"} and self.row is not None:
            self.cell = []
        elif tag in {"dt", "dd"}:
            self.definition = []

    def handle_data(self, text):
        if self.cell is not None:
            self.cell.append(text)
        if self.definition is not None:
            self.definition.append(text)

    def handle_endtag(self, tag):
        if tag in {"td", "th"} and self.cell is not None:
            self.row.append("".join(self.cell).strip())
            self.cell = None
        elif tag == "tr" and self.row is not None:
            self.rows.append(self.row)
            self.row = None
        elif tag in {"dt", "dd"} and self.definition is not None:
            text = "".join(self.definition).strip()
            if tag == "dt":
                self.term = text
            else:
                self.definitions[self.term] = text
            self.definition = None

    def product_rows(self):
        header = next(row for row in self.rows if "名称" in row)
        return {row[1]: dict(zip(header, row)) for row in self.rows[1:] if len(row) == len(header)}


def business_snapshot():
    with get_db() as conn:
        tables = [row[0] for row in conn.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name")]
        return {table: [tuple(row) for row in conn.execute(f'SELECT * FROM "{table}" ORDER BY 1')] for table in tables}


@pytest.fixture
def stock_products():
    app = create_app()
    cases = [
        ("C4虚构未初始化", None, None, "2", "未启用", "未启用"),
        ("C4虚构零库存", "0", "", "0", "已启用", "缺货"),
        ("C4虚构安全等号", "1.250", "3.333333", "1.250", "已启用", "库存告急"),
        ("C4虚构正常库存", "3", "3.333333", "2", "已启用", "正常"),
        ("C4虚构明确零成本", "2", "0", "0", "已启用", "正常"),
    ]
    products = []
    for index, (name, qty, cost, safety, enabled, status) in enumerate(cases):
        pid = create_product(name, "C4-B", "米" if index == 2 else "个", 1000, safety_stock=safety)
        if qty is not None:
            initialize_product(pid, qty, cost, date.today().isoformat(), "虚构测试期初", f"c4-b-{index}", confirm_zero=qty == "0")
        products.append({"id": pid, "name": name, "enabled": enabled, "status": status})
    return app.test_client(), products


def test_real_product_list_and_edit_show_independent_inventory_axes_without_writes(stock_products):
    client, products = stock_products
    before = business_snapshot()
    response = client.get("/products/")
    assert response.status_code == 200
    rows = ProductHTML(response.get_data(as_text=True)).product_rows()
    assert "启用状态" in rows[products[0]["name"]], "list needs independent enable column"
    assert "库存状况" in rows[products[0]["name"]], "list needs independent stock column"
    for product in products:
        row = rows[product["name"]]
        assert row["启用状态"] == product["enabled"]
        assert row["库存状况"] == product["status"]
        response = client.get(f'/products/{product["id"]}/edit')
        assert response.status_code == 200
        definitions = ProductHTML(response.get_data(as_text=True)).definitions
        assert definitions["启用状态"] == product["enabled"]
        assert definitions["库存状况"] == product["status"]
    assert business_snapshot() == before
    with get_db() as conn:
        assert conn.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
        assert conn.execute("PRAGMA foreign_key_check").fetchall() == []


def test_product_cost_display_distinguishes_empty_unknown_and_explicit_zero(stock_products):
    client, products = stock_products
    before = business_snapshot()
    rows = ProductHTML(client.get("/products/").get_data(as_text=True)).product_rows()
    expected = [
        ("—", "—", "—"),
        ("0", "无库存", "0.00"),
        ("1.25", "3.33", "4.17"),
        ("3", "3.33", "10.00"),
        ("2", "0.00", "0.00"),
    ]
    for product, (quantity, average, total) in zip(products, expected):
        row = rows[product["name"]]
        assert row["当前库存"] == quantity
        assert row["移动平均成本"] == average
        assert row["库存金额"] == total
        definitions = ProductHTML(client.get(f'/products/{product["id"]}/edit').get_data(as_text=True)).definitions
        unit = row["单位"]
        assert definitions[f"当前库存（{unit}）"] == quantity
        assert definitions[f"移动平均成本（元/{unit}）"] == average
        assert definitions["库存金额（元）"] == total
    # 3 x visible average 3.33 = 9.99, not the true total 10.00.
    assert business_snapshot() == before


def test_real_routes_keep_missing_average_and_disabled_inventory_unknown(stock_products):
    client, products = stock_products
    missing_avg, disabled = products[3], products[4]
    with get_db() as conn:
        # Nullable average is an explicitly supported schema state. Disabled
        # stored numbers must not leak as enabled current inventory or cost.
        conn.execute("UPDATE product_inventory_state SET avg_cost_micro=NULL WHERE product_id=?", (missing_avg["id"],))
        conn.execute("UPDATE product_inventory_state SET enabled=0 WHERE product_id=?", (disabled["id"],))
    before = business_snapshot()
    rows = ProductHTML(client.get("/products/").get_data(as_text=True)).product_rows()
    assert rows[missing_avg["name"]]["移动平均成本"] == "—"
    assert rows[missing_avg["name"]]["库存金额"] == "10.00"
    for field in ("当前库存", "移动平均成本", "库存金额"):
        assert rows[disabled["name"]][field] == "—"
    assert rows[disabled["name"]]["启用状态"] == "未启用"
    assert rows[disabled["name"]]["库存状况"] == "未启用"
    for product in (missing_avg, disabled):
        definitions = ProductHTML(client.get(f'/products/{product["id"]}/edit').get_data(as_text=True)).definitions
        assert definitions["移动平均成本（元/个）"] == "—"
    assert business_snapshot() == before


def test_unknown_quantity_template_projection_never_fabricates_quantity_or_average(stock_products):
    from flask import render_template
    from erp.routes.products import _product_list_context
    from erp.services.inventory_status import project_inventory_status
    from erp.utils.money import cents_to_yuan, micro_to_yuan

    client, products = stock_products
    with get_db() as conn:
        rows, suggestions, _truncated = _product_list_context(conn)
        state = dict(conn.execute("SELECT * FROM product_inventory_state WHERE product_id=?", (products[3]["id"],)).fetchone())
    product = next(row for row in rows if row["id"] == products[3]["id"])
    # Current schema forbids NULL quantity: test the pure/template defensive
    # seam without altering schema or inventing a persisted invalid state.
    product["quantity_3dp"] = state["quantity_3dp"] = None
    status = project_inventory_status(enabled=True, quantity_3dp=None, safety_stock_3dp=product["safety_stock_3dp"])
    product.update(status)
    before = business_snapshot()
    with client.application.test_request_context():
        list_html = render_template("products/list.html", products=[product], suggestions=suggestions, q="", cents_to_yuan=cents_to_yuan, micro_to_yuan=micro_to_yuan)
        edit_html = render_template("products/edit.html", product=product, inventory_state=state, inventory_status=status, initialization=None, cents_to_yuan=cents_to_yuan, micro_to_yuan=micro_to_yuan)
    row = ProductHTML(list_html).product_rows()[product["name"]]
    assert (row["启用状态"], row["库存状况"], row["当前库存"], row["移动平均成本"]) == ("已启用", "未知", "—", "—")
    definitions = ProductHTML(edit_html).definitions
    assert definitions["库存状况"] == "未知"
    assert definitions["当前库存（个）"] == "—"
    assert definitions["移动平均成本（元/个）"] == "—"
    assert business_snapshot() == before


def test_readonly_context_survives_search_export_and_real_error_renderers(stock_products):
    import io
    from openpyxl import load_workbook

    client, products = stock_products
    before = business_snapshot()
    responses = [
        (client.get("/products/", query_string={"q": "C4虚构"}), 200),
        (client.post("/products/import", data={}), 400),
        (client.post("/products/create", data={"name": "", "unit": "个", "default_price": "1", "safety_stock": "0"}), 400),
        (client.post(f'/products/{products[3]["id"]}/edit', data={"name": products[3]["name"], "spec": "C4-B", "unit": "箱", "default_price": "10", "safety_stock": "2"}), 400),
        (client.post(f'/products/{products[1]["id"]}/inventory/revise', data={"quantity": "1", "unit_cost": "1", "reason": "", "expected_version": "1"}), 400),
        (client.post(f'/products/{products[0]["id"]}/inventory/initialize', data={"quantity": "0", "unit_cost": "", "business_date": date.today().isoformat()}), 400),
        (client.post(f'/products/{products[3]["id"]}/image', data={}), 400),
    ]
    for response, expected_status in responses:
        assert response.status_code == expected_status
        parsed = ProductHTML(response.get_data(as_text=True))
        if parsed.rows:
            assert parsed.product_rows()[products[3]["name"]]["库存状况"] == "正常"
        else:
            assert "启用状态" in parsed.definitions and "库存状况" in parsed.definitions
    export = client.get("/products/export.xlsx", query_string={"q": "C4虚构"})
    assert export.status_code == 200
    workbook = load_workbook(io.BytesIO(export.data))
    assert workbook.active.max_row == 6
    assert {row[0] for row in workbook.active.iter_rows(min_row=2, values_only=True)} == {p["name"] for p in products}
    assert business_snapshot() == before


def test_successful_import_uses_same_two_axis_read_projection(stock_products):
    import io

    client, products = stock_products
    response = client.post("/products/import", data={"file": (io.BytesIO("商品名称,型号,价格\nC4虚构导入新品,X,1.00\n".encode("utf-8-sig")), "synthetic.csv")}, content_type="multipart/form-data")
    assert response.status_code == 200
    rows = ProductHTML(response.get_data(as_text=True)).product_rows()
    assert rows["C4虚构导入新品"]["启用状态"] == "未启用"
    assert rows["C4虚构导入新品"]["库存状况"] == "未启用"
    assert rows[products[1]["name"]]["库存状况"] == "缺货"
    with get_db() as conn:
        assert conn.execute("SELECT COUNT(*) FROM product_inventory_state").fetchone()[0] == 4


def test_workbench_counts_whole_enabled_alert_scope_not_limited_rank(stock_products):
    client, products = stock_products
    for index in range(11):
        pid = create_product(f"C4虚构额外提醒{index:02}", "W", "个", 100, safety_stock="1")
        initialize_product(pid, "1", "0", date.today().isoformat(), "虚构", f"c4-b-alert-{index}")
    deleted = create_product("C4虚构已删除提醒", "W", "个", 100)
    initialize_product(deleted, "0", "", date.today().isoformat(), "虚构", "c4-b-alert-deleted", confirm_zero=True)
    disabled = create_product("C4虚构未启用提醒", "W", "个", 100)
    with get_db() as conn:
        conn.execute("UPDATE products SET deleted_at=CURRENT_TIMESTAMP WHERE id=?", (deleted,))
        conn.execute("INSERT INTO product_inventory_state(product_id, enabled, quantity_3dp) VALUES (?, 0, 0)", (disabled,))
    before = business_snapshot()
    html = client.get("/").get_data(as_text=True)
    alert = html.split('id="dashboardAlerts"', 1)[1].split("</section>", 1)[0]
    assert "13 个商品需要关注" in alert
    assert business_snapshot() == before

